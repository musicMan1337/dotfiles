#!/usr/bin/env python3
"""
Budget regeneration script.
Reads a JSON budget model from stdin, writes the formatted Excel file and markdown.

Usage:
    echo '{"salary": 107000, ...}' | python3 budget-regen.py

The JSON model schema:
{
    "salary": 107000,
    "pretax": [{"name": "401k (6.0%)", "amount": 267.50}, ...],
    "aftertax": [{"name": "Accident Insurance", "amount": 8.07}, ...],
    "taxes": [{"name": "Federal", "pct": 0.168, "amount": 750.75}, ...],
    "ally_deposit": 1820.00,
    "doge_deposit": 150.00,
    "ally": {
        "housing": [{"name": "Rent", "monthly": 1362.00, "notes": "Lion contributes $800/mo"}, ...],
        "auto": [{"name": "Kylo Ren (Mazda 3)", "monthly": 387.50}, ...],
        "savings": [{"name": "Presents", "monthly": 100.00}, ...]
    },
    "lion": [{"name": "Gas", "monthly": 100.00}, ...],
    "subscriptions": [{"name": "Costco", "monthly": 10.00}, ...],
    "doge": [{"name": "Pet Insurance", "monthly": 48.46}, ...],
    "wellsfargo": {
        "fixed": [{"name": "Groceries", "monthly": 500.00}, ...],
        "discretionary": [{"name": "Travel", "monthly": 0}, ...]
    }
}
"""

import json
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

EXCEL_PATH = "/Users/derek/eBacon/obsidian/eBacon/finances/boo boo budget.xlsx"
MD_PATH = "/Users/derek/eBacon/obsidian/eBacon/finances/Boo Boo Budget.md"

# ── Styles ──────────────────────────────────────────────────
bold = Font(bold=True)
bold_white = Font(bold=True, color="FFFFFF", size=11)
italic = Font(italic=True, size=9, color="666666")
currency = '"$"#,##0.00'
pct_fmt = '0.0%'

income_header = PatternFill("solid", fgColor="1F3864")
ally_header = PatternFill("solid", fgColor="375623")
lion_header = PatternFill("solid", fgColor="8B0000")
doge_header = PatternFill("solid", fgColor="7B3F00")
regular_header = PatternFill("solid", fgColor="4A1A6B")

income_section = PatternFill("solid", fgColor="2F5496")
ally_section = PatternFill("solid", fgColor="548235")
lion_section = PatternFill("solid", fgColor="CC4444")
doge_section = PatternFill("solid", fgColor="BF6900")
regular_section = PatternFill("solid", fgColor="7030A0")

income_colhead = PatternFill("solid", fgColor="B4C6E7")
ally_colhead = PatternFill("solid", fgColor="C6E0B4")
lion_colhead = PatternFill("solid", fgColor="F4CCCC")
doge_colhead = PatternFill("solid", fgColor="F9CB9C")
regular_colhead = PatternFill("solid", fgColor="D9C2EC")

income_stripe = PatternFill("solid", fgColor="D6E4F0")
ally_stripe = PatternFill("solid", fgColor="E2EFDA")
lion_stripe = PatternFill("solid", fgColor="FCE4EC")
doge_stripe = PatternFill("solid", fgColor="FFF3E0")
regular_stripe = PatternFill("solid", fgColor="F3E5F5")

total_fill = PatternFill("solid", fgColor="DDDDDD")
gold_fill = PatternFill("solid", fgColor="FFC000")
arrow_fill = PatternFill("solid", fgColor="E8E8E8")

thin_bottom = Border(bottom=Side(style='thin', color='CCCCCC'))
thick = Side(style='thick', color='333333')

CL = openpyxl.utils.get_column_letter


# ── Helpers ─────────────────────────────────────────────────
def apply_thick_borders(ws, regions):
    for r1, r2, c1, c2 in regions:
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                cell = ws.cell(row=r, column=c)
                ex = cell.border
                cell.border = Border(
                    top=thick if r == r1 else ex.top,
                    bottom=thick if r == r2 else ex.bottom,
                    left=thick if c == c1 else ex.left,
                    right=thick if c == c2 else ex.right,
                )


def write_account_header(ws, row, col, text, fill):
    for r in range(row, row + 2):
        for c in range(col, col + 4):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill
            cell.font = Font(bold=True, color="FFFFFF", size=13)
            cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=row, column=col, value=text)
    ws.merge_cells(start_row=row, start_column=col, end_row=row + 1, end_column=col + 3)
    return row + 2


def write_section_header(ws, row, col, text, fill):
    for c in range(col, col + 4):
        ws.cell(row=row, column=c).fill = fill
        ws.cell(row=row, column=c).font = bold_white
    ws.cell(row=row, column=col, value=text)
    return row + 1


def write_col_headers(ws, row, col, headers, fill):
    for i, h in enumerate(headers):
        cell = ws.cell(row=row, column=col + i, value=h)
        cell.font = Font(bold=True, size=10)
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center' if i > 0 else 'left')
        cell.border = Border(bottom=Side(style='thin', color='AAAAAA'))
    return row + 1


def style_row(ws, row, col, is_total=False, stripe=None):
    if is_total:
        for c in range(col, col + 4):
            ws.cell(row=row, column=c).fill = total_fill
            ws.cell(row=row, column=c).font = bold
            ws.cell(row=row, column=c).border = Border(
                top=Side(style='thin', color='999999'),
                bottom=Side(style='medium', color='999999'),
            )
    elif stripe:
        for c in range(col, col + 4):
            ws.cell(row=row, column=c).fill = stripe
            ws.cell(row=row, column=c).border = thin_bottom
    else:
        for c in range(col, col + 4):
            ws.cell(row=row, column=c).border = thin_bottom


def write_item_monthly(ws, row, col, label, monthly_val, notes=None, stripe=None):
    style_row(ws, row, col, stripe=stripe)
    ws.cell(row=row, column=col, value=label)
    ws.cell(row=row, column=col + 1, value=monthly_val).number_format = currency
    ws.cell(row=row, column=col + 2).value = f"={CL(col+1)}{row}/2"
    ws.cell(row=row, column=col + 2).number_format = currency
    if notes:
        ws.cell(row=row, column=col + 3, value=notes).font = Font(size=9, color="666666")
    return row + 1


def write_item_pc_only(ws, row, col, label, pc_val, notes=None, stripe=None):
    style_row(ws, row, col, stripe=stripe)
    ws.cell(row=row, column=col, value=label)
    ws.cell(row=row, column=col + 2, value=pc_val).number_format = currency
    if notes:
        ws.cell(row=row, column=col + 3, value=notes).font = Font(size=9, color="666666")
    return row + 1


def write_sum_total(ws, row, col, label, item_rows, has_monthly=True):
    style_row(ws, row, col, is_total=True)
    ws.cell(row=row, column=col, value=label)
    pc = CL(col + 2)
    ws.cell(row=row, column=col + 2).value = "=" + "+".join(f"{pc}{r}" for r in item_rows)
    ws.cell(row=row, column=col + 2).number_format = currency
    ws.cell(row=row, column=col + 2).font = bold
    if has_monthly:
        mo = CL(col + 1)
        ws.cell(row=row, column=col + 1).value = "=" + "+".join(f"{mo}{r}" for r in item_rows)
        ws.cell(row=row, column=col + 1).number_format = currency
        ws.cell(row=row, column=col + 1).font = bold
    return row + 1


def write_grand_total(ws, row, col, label, total_rows, fill, has_monthly=True):
    for r in range(row, row + 2):
        for c in range(col, col + 4):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(vertical='center')
    ws.cell(row=row, column=col, value=label)
    ws.merge_cells(start_row=row, start_column=col, end_row=row + 1, end_column=col)
    pc = CL(col + 2)
    ws.cell(row=row, column=col + 2).value = "=" + "+".join(f"{pc}{r}" for r in total_rows)
    ws.cell(row=row, column=col + 2).number_format = currency
    ws.cell(row=row, column=col + 2).font = Font(bold=True, color="FFFFFF", size=11)
    ws.merge_cells(start_row=row, start_column=col + 2, end_row=row + 1, end_column=col + 2)
    if has_monthly:
        mo = CL(col + 1)
        ws.cell(row=row, column=col + 1).value = "=" + "+".join(f"{mo}{r}" for r in total_rows)
        ws.cell(row=row, column=col + 1).number_format = currency
        ws.cell(row=row, column=col + 1).font = Font(bold=True, color="FFFFFF", size=11)
    ws.merge_cells(start_row=row, start_column=col + 1, end_row=row + 1, end_column=col + 1)
    ws.merge_cells(start_row=row, start_column=col + 3, end_row=row + 1, end_column=col + 3)
    return row + 2


# ── Main ────────────────────────────────────────────────────
def generate_excel(data):
    wb = openpyxl.load_workbook(EXCEL_PATH)
    if "Budget (Updated)" in wb.sheetnames:
        del wb["Budget (Updated)"]
    ws = wb.create_sheet("Budget (Updated)", 0)

    for col, w in {'A': 28, 'B': 14, 'C': 14, 'D': 24, 'E': 3, 'F': 28, 'G': 14, 'H': 14, 'I': 24, 'J': 3, 'K': 28, 'L': 14, 'M': 14, 'N': 24}.items():
        ws.column_dimensions[col].width = w

    table_regions = []
    keys = {}

    # ── INCOME (A-D) ────────────────────────────────────────
    c = 1; r = 1; ts = r
    r = write_account_header(ws, r, c, "INCOME", income_header)
    r = write_section_header(ws, r, c, "Gross Pay", income_section)
    r = write_col_headers(ws, r, c, ["", "Annual", "Per Check", ""], income_colhead)
    style_row(ws, r, c, is_total=True)
    ws.cell(row=r, column=c, value="Salary")
    ws.cell(row=r, column=c + 1, value=data["salary"]).number_format = currency
    ws.cell(row=r, column=c + 1).font = bold
    ws.cell(row=r, column=c + 2).value = f"=B{r}/24"
    ws.cell(row=r, column=c + 2).number_format = currency
    ws.cell(row=r, column=c + 2).font = bold
    keys['gross'] = r; r += 2

    r = write_section_header(ws, r, c, "Pre-Tax Deductions", income_section)
    r = write_col_headers(ws, r, c, ["Deduction", "", "Per Check", ""], income_colhead)
    pt_rows = []
    for i, item in enumerate(data["pretax"]):
        r = write_item_pc_only(ws, r, c, item["name"], item["amount"], stripe=income_stripe if i % 2 == 0 else None)
        pt_rows.append(r - 1)
    r = write_sum_total(ws, r, c, "Total Pre-Tax", pt_rows, has_monthly=False)
    keys['pretax'] = r - 1; r += 1

    r = write_section_header(ws, r, c, "After-Tax Deductions", income_section)
    r = write_col_headers(ws, r, c, ["Deduction", "", "Per Check", ""], income_colhead)
    at_rows = []
    for i, item in enumerate(data["aftertax"]):
        r = write_item_pc_only(ws, r, c, item["name"], item["amount"], stripe=income_stripe if i % 2 == 0 else None)
        at_rows.append(r - 1)
    r = write_sum_total(ws, r, c, "Total After-Tax", at_rows, has_monthly=False)
    keys['aftertax'] = r - 1; r += 1

    r = write_section_header(ws, r, c, "Taxes", income_section)
    r = write_col_headers(ws, r, c, ["Tax", "% of Gross", "Per Check", ""], income_colhead)
    tx_rows = []
    for i, item in enumerate(data["taxes"]):
        stripe = income_stripe if i % 2 == 0 else None
        if stripe:
            for cc in range(c, c + 4): ws.cell(row=r, column=cc).fill = stripe; ws.cell(row=r, column=cc).border = thin_bottom
        else:
            for cc in range(c, c + 4): ws.cell(row=r, column=cc).border = thin_bottom
        ws.cell(row=r, column=c, value=item["name"])
        ws.cell(row=r, column=c + 1, value=item["pct"]).number_format = pct_fmt
        ws.cell(row=r, column=c + 2, value=item["amount"]).number_format = currency
        tx_rows.append(r); r += 1
    style_row(ws, r, c, is_total=True)
    ws.cell(row=r, column=c, value="Total Taxes")
    ws.cell(row=r, column=c + 2).value = "=" + "+".join(f"C{tr}" for tr in tx_rows)
    ws.cell(row=r, column=c + 2).number_format = currency; ws.cell(row=r, column=c + 2).font = bold
    ws.cell(row=r, column=c + 1).value = f"=C{r}/C{keys['gross']}"
    ws.cell(row=r, column=c + 1).number_format = pct_fmt; ws.cell(row=r, column=c + 1).font = bold
    keys['taxes'] = r; r += 2

    r = write_section_header(ws, r, c, "Direct Deposit Split", income_section)
    r = write_col_headers(ws, r, c, ["Account", "", "Per Check", ""], income_colhead)
    r = write_item_pc_only(ws, r, c, "Ally", data["ally_deposit"])
    keys['ally_dep'] = r - 1
    r = write_item_pc_only(ws, r, c, "Doge", data["doge_deposit"], stripe=income_stripe)
    keys['doge_dep'] = r - 1
    style_row(ws, r, c)
    ws.cell(row=r, column=c, value="Wells Fargo")
    ws.cell(row=r, column=c + 2).value = f"=(C{keys['gross']}-C{keys['pretax']}-C{keys['aftertax']}-C{keys['taxes']})-C{keys['ally_dep']}-C{keys['doge_dep']}"
    ws.cell(row=r, column=c + 2).number_format = currency
    keys['wf_dep'] = r; r += 1

    for rr in range(r, r + 2):
        for cc in range(c, c + 4):
            cell = ws.cell(row=rr, column=cc)
            cell.fill = income_header; cell.font = Font(bold=True, color="FFFFFF", size=11); cell.alignment = Alignment(vertical='center')
    ws.cell(row=r, column=c, value="TOTAL TAKE-HOME")
    ws.merge_cells(start_row=r, start_column=c, end_row=r + 1, end_column=c)
    ws.cell(row=r, column=c + 2).value = f"=C{keys['gross']}-C{keys['pretax']}-C{keys['aftertax']}-C{keys['taxes']}"
    ws.cell(row=r, column=c + 2).number_format = currency
    ws.cell(row=r, column=c + 2).font = Font(bold=True, color="FFFFFF", size=11)
    ws.merge_cells(start_row=r, start_column=c + 1, end_row=r + 1, end_column=c + 1)
    ws.merge_cells(start_row=r, start_column=c + 2, end_row=r + 1, end_column=c + 2)
    ws.merge_cells(start_row=r, start_column=c + 3, end_row=r + 1, end_column=c + 3)
    keys['net'] = r; r += 2
    table_regions.append((ts, r - 1, c, c + 3))

    # ── ALLY + LION (F-I) ──────────────────────────────────
    c = 6; r = 1; ts = r
    r = write_account_header(ws, r, c, f"ALLY \u2014 ${data['ally_deposit']:,.2f}/check", ally_header)
    ws.cell(row=r, column=c, value="Sub-accounts handle distribution via internal recurring transfers.").font = italic; r += 1

    def write_ally_section(section_name, section_label, items, section_fill, colhead_fill, stripe_fill):
        nonlocal r
        r = write_section_header(ws, r, c, section_label, section_fill)
        hdrs = ["Expense", "Monthly", "Per Check", "Notes"] if section_name != "savings" else ["Item", "Monthly", "Per Check", "Notes"]
        r = write_col_headers(ws, r, c, hdrs, colhead_fill)
        rows = []
        for i, item in enumerate(items):
            r = write_item_monthly(ws, r, c, item["name"], item["monthly"], notes=item.get("notes"), stripe=stripe_fill if i % 2 == 0 else None)
            rows.append(r - 1)
        r = write_sum_total(ws, r, c, f"Total {section_label}", rows)
        return r - 1, rows

    housing_total, _ = write_ally_section("housing", "Housing", data["ally"]["housing"], ally_section, ally_colhead, ally_stripe); r += 1
    auto_total, _ = write_ally_section("auto", "Auto & Debt", data["ally"]["auto"], ally_section, ally_colhead, ally_stripe); r += 1
    savings_total, _ = write_ally_section("savings", "Savings", data["ally"]["savings"], ally_section, ally_colhead, ally_stripe)
    r = write_grand_total(ws, r, c, "TOTAL ALLY", [housing_total, auto_total, savings_total], ally_header)

    # Arrow row
    for cc in range(c, c + 4):
        cell = ws.cell(row=r, column=cc)
        cell.fill = arrow_fill; cell.font = Font(size=14, color="555555")
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.value = "\u25BC"
    r += 1

    # Lion
    r = write_account_header(ws, r, c, "LION (cabanilla.km@gmail.com)", lion_header)
    ws.cell(row=r, column=c, value="Sub-account within Ally. Transferred internally \u2014 tracked for visibility.").font = italic; r += 1
    r = write_section_header(ws, r, c, "Lion Expenses", lion_section)
    r = write_col_headers(ws, r, c, ["Expense", "Monthly", "Per Check", "Notes"], lion_colhead)
    lion_rows = []
    for i, item in enumerate(data["lion"]):
        r = write_item_monthly(ws, r, c, item["name"], item["monthly"], notes=item.get("notes"), stripe=lion_stripe if i % 2 == 0 else None)
        lion_rows.append(r - 1)
    r = write_grand_total(ws, r, c, "TOTAL LION", lion_rows, lion_header)
    r += 1

    r = write_section_header(ws, r, c, "Subscriptions Detail", lion_section)
    r = write_col_headers(ws, r, c, ["Service", "Monthly", "", ""], lion_colhead)
    sub_rows = []
    for i, item in enumerate(data["subscriptions"]):
        style_row(ws, r, c, stripe=lion_stripe if i % 2 == 0 else None)
        ws.cell(row=r, column=c, value=item["name"])
        ws.cell(row=r, column=c + 1, value=item["monthly"]).number_format = currency
        sub_rows.append(r); r += 1
    r = write_sum_total(ws, r, c, "Total Subscriptions", sub_rows, has_monthly=True)
    ws.cell(row=r - 1, column=c + 2).value = None
    table_regions.append((ts, r - 1, c, c + 3))

    # ── DOGE (K-N) ─────────────────────────────────────────
    c = 11; r = 1; ts = r
    r = write_account_header(ws, r, c, f"DOGE \u2014 ${data['doge_deposit']:,.2f}/check", doge_header)
    ws.cell(row=r, column=c, value="Single account, no sub-transfers.").font = italic; r += 1
    r = write_section_header(ws, r, c, "Pet Expenses", doge_section)
    r = write_col_headers(ws, r, c, ["Expense", "Monthly", "Per Check", "Notes"], doge_colhead)
    doge_rows = []
    for i, item in enumerate(data["doge"]):
        r = write_item_monthly(ws, r, c, item["name"], item["monthly"], notes=item.get("notes"), stripe=doge_stripe if i % 2 == 0 else None)
        doge_rows.append(r - 1)
    r = write_grand_total(ws, r, c, "TOTAL DOGE", doge_rows, doge_header)
    table_regions.append((ts, r - 1, c, c + 3))
    r += 6

    # ── WELLS FARGO (K-N, below Doge) ──────────────────────
    ts = r
    r = write_account_header(ws, r, c, "WELLS FARGO", regular_header)
    ws.cell(row=r, column=c, value="Remainder after Ally & Doge. Day-to-day spending.").font = italic; r += 1

    r = write_section_header(ws, r, c, "Fixed", regular_section)
    r = write_col_headers(ws, r, c, ["Expense", "Monthly", "Per Check", "Notes"], regular_colhead)
    fix_rows = []
    for i, item in enumerate(data["wellsfargo"]["fixed"]):
        r = write_item_monthly(ws, r, c, item["name"], item["monthly"], notes=item.get("notes"), stripe=regular_stripe if i % 2 == 0 else None)
        fix_rows.append(r - 1)
    r = write_sum_total(ws, r, c, "Total Fixed", fix_rows)
    keys['wf_fixed'] = r - 1; r += 1

    r = write_section_header(ws, r, c, "Discretionary", regular_section)
    r = write_col_headers(ws, r, c, ["Category", "Monthly", "Per Check", "Notes"], regular_colhead)
    disc_rows = []
    for i, item in enumerate(data["wellsfargo"]["discretionary"]):
        r = write_item_monthly(ws, r, c, item["name"], item["monthly"], notes=item.get("notes"), stripe=regular_stripe if i % 2 == 0 else None)
        disc_rows.append(r - 1)
    r = write_sum_total(ws, r, c, "Total Discretionary", disc_rows)
    keys['wf_disc'] = r - 1; r += 1

    # Unallocated
    pc_col = CL(c + 2)
    for rr in range(r, r + 2):
        for cc in range(c, c + 4):
            cell = ws.cell(row=rr, column=cc)
            cell.fill = gold_fill; cell.font = Font(bold=True, size=13); cell.alignment = Alignment(vertical='center')
    ws.cell(row=r, column=c, value="UNALLOCATED")
    ws.merge_cells(start_row=r, start_column=c, end_row=r + 1, end_column=c)
    ws.cell(row=r, column=c + 2).value = f"=(C{keys['net']}-C{keys['ally_dep']}-C{keys['doge_dep']})-{pc_col}{keys['wf_fixed']}-{pc_col}{keys['wf_disc']}"
    ws.cell(row=r, column=c + 2).number_format = currency
    ws.cell(row=r, column=c + 2).font = Font(bold=True, size=13)
    ws.merge_cells(start_row=r, start_column=c + 1, end_row=r + 1, end_column=c + 1)
    ws.merge_cells(start_row=r, start_column=c + 2, end_row=r + 1, end_column=c + 2)
    ws.merge_cells(start_row=r, start_column=c + 3, end_row=r + 1, end_column=c + 3)
    table_regions.append((ts, r + 1, c, c + 3))

    apply_thick_borders(ws, table_regions)
    ws.freeze_panes = 'A3'
    wb.save(EXCEL_PATH)
    return keys


def generate_markdown(data):
    """Generate the markdown budget file from the data model."""
    lines = []
    a = lines.append

    def fmt(v):
        if v < 0:
            return f"-${abs(v):,.2f}"
        return f"${v:,.2f}"

    def pct(v):
        return f"{v*100:.1f}%"

    salary = data["salary"]
    pc = salary / 24
    pretax_total = sum(i["amount"] for i in data["pretax"])
    aftertax_total = sum(i["amount"] for i in data["aftertax"])
    tax_total = sum(i["amount"] for i in data["taxes"])
    net = pc - pretax_total - aftertax_total - tax_total
    wf_deposit = net - data["ally_deposit"] - data["doge_deposit"]

    a("# Boo Boo Budget")
    a("")
    a("> Original spreadsheet: [[boo boo budget.xlsx]]")
    a("")
    a("---")
    a("")
    a("## Income")
    a("")
    a("### Gross Pay")
    a("| | Per Check | Annual |")
    a("|---|---:|---:|")
    a(f"| **Gross** | {fmt(pc)} | {fmt(salary)} |")
    a("")

    a("### Pre-Tax Deductions")
    a("| Deduction | Per Check |")
    a("|---|---:|")
    for i in data["pretax"]:
        a(f"| {i['name']} | {fmt(i['amount'])} |")
    a(f"| **Total Pre-Tax** | **{fmt(pretax_total)}** |")
    a("")

    a("### After-Tax Deductions")
    a("| Deduction | Per Check |")
    a("|---|---:|")
    for i in data["aftertax"]:
        a(f"| {i['name']} | {fmt(i['amount'])} |")
    a(f"| **Total After-Tax** | **{fmt(aftertax_total)}** |")
    a("")

    tax_pct_total = tax_total / pc if pc else 0
    a(f"### Taxes ({pct(tax_pct_total)} of gross)")
    a("| Tax | Per Check | % of Gross |")
    a("|---|---:|---:|")
    for i in data["taxes"]:
        a(f"| {i['name']} | {fmt(i['amount'])} | {pct(i['pct'])} |")
    a(f"| **Total** | **{fmt(tax_total)}** | **{pct(tax_pct_total)}** |")
    a("")

    a("### Net Pay & Direct Deposit Split")
    a("| Account | Per Check |")
    a("|---|---:|")
    a(f"| Ally (mammoth17@gmail.com) | {fmt(data['ally_deposit'])} |")
    a(f"| Doge | {fmt(data['doge_deposit'])} |")
    a(f"| Wells Fargo | {fmt(wf_deposit)} |")
    a(f"| **Total Take-Home** | **{fmt(net)}** |")
    a("")
    a("---")
    a("")

    # Ally
    a(f"## Ally Account \u2014 {fmt(data['ally_deposit'])}/check")
    a("")
    a("> Sub-accounts handle distribution via internal recurring transfers.")
    a("")

    for section_key, section_label in [("housing", "Housing"), ("auto", "Auto & Debt"), ("savings", "Savings")]:
        items = data["ally"][section_key]
        total = sum(i["monthly"] for i in items)
        a(f"### {section_label} \u2014 {fmt(total)}/mo")
        has_notes = any(i.get("notes") for i in items)
        if has_notes:
            a("| Expense | Monthly | Per Check | Notes |")
            a("|---|---:|---:|---|")
            for i in items:
                a(f"| {i['name']} | {fmt(i['monthly'])} | {fmt(i['monthly']/2)} | {i.get('notes', '')} |")
        else:
            a("| Expense | Monthly | Per Check |")
            a("|---|---:|---:|")
            for i in items:
                a(f"| {i['name']} | {fmt(i['monthly'])} | {fmt(i['monthly']/2)} |")
        a("")

    lion_total = sum(i["monthly"] for i in data["lion"])
    a(f"### Lion (cabanilla.km@gmail.com) \u2014 {fmt(lion_total)}/mo")
    a("> Sub-account within Ally. Transferred internally \u2014 tracked here for visibility.")
    a("")
    a("| Expense | Monthly | Per Check |")
    a("|---|---:|---:|")
    for i in data["lion"]:
        a(f"| {i['name']} | {fmt(i['monthly'])} | {fmt(i['monthly']/2)} |")
    a("")

    sub_total = sum(i["monthly"] for i in data["subscriptions"])
    a(f"#### Subscriptions Detail \u2014 {fmt(sub_total)}/mo")
    a("| Service | Monthly |")
    a("|---|---:|")
    for i in data["subscriptions"]:
        a(f"| {i['name']} | {fmt(i['monthly'])} |")
    a("")
    a("---")
    a("")

    # Doge
    doge_total = sum(i["monthly"] for i in data["doge"])
    a(f"## Doge Account \u2014 {fmt(data['doge_deposit'])}/check")
    a("")
    a("> Single account, no sub-transfers.")
    a("")
    a("| Expense | Monthly | Per Check |")
    a("|---|---:|---:|")
    for i in data["doge"]:
        a(f"| {i['name']} | {fmt(i['monthly'])} | {fmt(i['monthly']/2)} |")
    a(f"| **Total** | **{fmt(doge_total)}** | **{fmt(doge_total/2)}** |")
    a("")
    a("---")
    a("")

    # Wells Fargo
    a(f"## Wells Fargo \u2014 {fmt(wf_deposit)}/check")
    a("")
    a("> Remainder after Ally & Doge direct deposits. Day-to-day spending.")
    a("")

    for section_key, section_label in [("fixed", "Fixed"), ("discretionary", "Discretionary")]:
        items = data["wellsfargo"][section_key]
        a(f"### {section_label}")
        a("| Expense | Monthly | Per Check |")
        a("|---|---:|---:|")
        for i in items:
            a(f"| {i['name']} | {fmt(i['monthly'])} | {fmt(i['monthly']/2)} |")
        a("")

    wf_fixed_total = sum(i["monthly"] for i in data["wellsfargo"]["fixed"]) / 2
    wf_disc_total = sum(i["monthly"] for i in data["wellsfargo"]["discretionary"]) / 2
    unallocated = wf_deposit - wf_fixed_total - wf_disc_total

    a("### Remaining After All Expenses")
    a("| | Per Check |")
    a("|---|---:|")
    a(f"| **Unallocated** | **{fmt(unallocated)}** |")
    a("")

    with open(MD_PATH, "w") as f:
        f.write("\n".join(lines))


if __name__ == "__main__":
    data = json.load(sys.stdin)
    generate_excel(data)
    generate_markdown(data)
    print(json.dumps({"status": "ok", "excel": EXCEL_PATH, "markdown": MD_PATH}))
