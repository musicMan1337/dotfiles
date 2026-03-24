#!/usr/bin/env python3
"""
Read current budget state from Excel and output as JSON.
Used by the budget-adjust skill to understand current state before modifications.

Usage:
    python3 budget-read.py
"""

import json
import openpyxl

EXCEL_PATH = "/Users/derek/eBacon/obsidian/eBacon/finances/boo boo budget.xlsx"


def read_budget():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["Budget (Updated)"]

    def val(row, col):
        v = ws.cell(row=row, column=col).value
        return v if v is not None else 0

    def sval(row, col):
        v = ws.cell(row=row, column=col).value
        return str(v) if v else ""

    # Walk through the sheet to find sections dynamically
    # Income is in columns A-D (1-4)
    # Ally+Lion is in columns F-I (6-9)
    # Doge+WF is in columns K-N (11-14)

    data = {
        "salary": None,
        "pretax": [],
        "aftertax": [],
        "taxes": [],
        "ally_deposit": None,
        "doge_deposit": None,
        "ally": {"housing": [], "auto": [], "savings": []},
        "lion": [],
        "subscriptions": [],
        "doge": [],
        "wellsfargo": {"fixed": [], "discretionary": []}
    }

    # Parse Income column (A-D)
    section = None
    for r in range(1, 60):
        label = sval(r, 1).strip()
        if label == "Salary":
            data["salary"] = val(r, 2)  # Annual is in col B
            if not data["salary"] or data["salary"] == 0:
                data["salary"] = val(r, 3) * 24  # Fallback: per-check * 24
        elif label == "Pre-Tax Deductions":
            section = "pretax"
        elif label == "After-Tax Deductions":
            section = "aftertax"
        elif label == "Taxes":
            section = "taxes"
        elif label == "Direct Deposit Split":
            section = "deposits"
        elif label.startswith("Total") or label.startswith("TOTAL"):
            continue
        elif section == "pretax" and label and label != "Deduction":
            data["pretax"].append({"name": label, "amount": val(r, 3)})
        elif section == "aftertax" and label and label != "Deduction":
            data["aftertax"].append({"name": label, "amount": val(r, 3)})
        elif section == "taxes" and label and label != "Tax":
            pct = val(r, 2)
            data["taxes"].append({"name": label, "pct": pct if pct else 0, "amount": val(r, 3)})
        elif section == "deposits":
            if "Ally" in label:
                data["ally_deposit"] = val(r, 3)
            elif "Doge" in label:
                data["doge_deposit"] = val(r, 3)

    # Parse Ally column (F-I)
    section = None
    for r in range(1, 80):
        label = sval(r, 6).strip()
        monthly = val(r, 7)
        notes = sval(r, 9)
        if label == "Housing":
            section = "housing"
        elif label == "Auto & Debt":
            section = "auto"
        elif label == "Savings":
            section = "savings"
        elif label == "Lion Expenses":
            section = "lion"
        elif label == "Subscriptions Detail":
            section = "subs"
        elif label.startswith("Total") or label.startswith("TOTAL") or label.startswith("\u25BC"):
            continue
        elif label in ("Expense", "Item", "Service", "Category", ""):
            continue
        elif "Sub-account" in label or "Sub-accounts" in label:
            continue
        elif section == "housing" and monthly:
            item = {"name": label, "monthly": monthly}
            if notes: item["notes"] = notes
            data["ally"]["housing"].append(item)
        elif section == "auto" and monthly:
            item = {"name": label, "monthly": monthly}
            if notes: item["notes"] = notes
            data["ally"]["auto"].append(item)
        elif section == "savings" and monthly:
            item = {"name": label, "monthly": monthly}
            if notes: item["notes"] = notes
            data["ally"]["savings"].append(item)
        elif section == "lion" and monthly:
            item = {"name": label, "monthly": monthly}
            if notes: item["notes"] = notes
            data["lion"].append(item)
        elif section == "subs" and monthly:
            data["subscriptions"].append({"name": label, "monthly": monthly})

    # Parse Doge + WF column (K-N)
    section = None
    for r in range(1, 80):
        label = sval(r, 11).strip()
        monthly = val(r, 12)
        notes = sval(r, 14)
        if label == "Pet Expenses":
            section = "doge"
        elif "WELLS FARGO" in label.upper():
            section = "wf_header"
        elif label == "Fixed":
            section = "wf_fixed"
        elif label == "Discretionary":
            section = "wf_disc"
        elif label.startswith("Total") or label.startswith("TOTAL") or label == "UNALLOCATED":
            continue
        elif label in ("Expense", "Category", ""):
            continue
        elif "Remainder" in label or "Single account" in label:
            continue
        elif section == "doge" and label:
            item = {"name": label, "monthly": monthly if monthly else 0}
            if notes: item["notes"] = notes
            data["doge"].append(item)
        elif section == "wf_fixed" and label:
            item = {"name": label, "monthly": monthly if monthly else 0}
            if notes: item["notes"] = notes
            data["wellsfargo"]["fixed"].append(item)
        elif section == "wf_disc" and label:
            item = {"name": label, "monthly": monthly if monthly else 0}
            if notes: item["notes"] = notes
            data["wellsfargo"]["discretionary"].append(item)

    return data


if __name__ == "__main__":
    data = read_budget()
    print(json.dumps(data, indent=2))
